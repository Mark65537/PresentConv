
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


def set_run_style(run, size=20, bold=False, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(40, 45, 55)


def add_markdown_text(paragraph, text, size=20):
    """Обрабатывает простой Markdown: жирный и курсивный текст."""
    pattern = r"(\*\*.*?\*\*|__.*?__|\*.*?\*|_.*?_)"

    for part in re.split(pattern, text):
        if not part:
            continue

        if (part.startswith("**") and part.endswith("**")
                or part.startswith("__") and part.endswith("__")):
            run = paragraph.add_run()
            run.text = part[2:-2]
            set_run_style(run, size, bold=True)

        elif (part.startswith("*") and part.endswith("*")
              or part.startswith("_") and part.endswith("_")):
            run = paragraph.add_run()
            run.text = part[1:-1]
            set_run_style(run, size, italic=True)

        else:
            run = paragraph.add_run()
            run.text = part
            set_run_style(run, size)


def add_slide(prs, title, lines, title_slide=False):
    layout = prs.slide_layouts[0] if title_slide else prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)

    if title_slide:
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(34)

        subtitle = slide.placeholders[1]
        subtitle.text = "\n".join(lines) if lines else ""
        for paragraph in subtitle.text_frame.paragraphs:
            for run in paragraph.runs:
                set_run_style(run, 18)
        return

    # Заголовок слайда
    title_box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8)
    )
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.name = "Arial"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 75, 130)

    # Основное содержимое
    body = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.4)
    )
    frame = body.text_frame
    frame.word_wrap = True
    frame.margin_left = Inches(0.1)
    frame.margin_right = Inches(0.1)
    frame.vertical_anchor = MSO_ANCHOR.TOP

    for line in lines:
        if not line.strip():
            continue

        if line.startswith("```"):
            continue

        p = frame.paragraphs[0] if not frame.paragraphs[0].text else frame.add_paragraph()

        if line.startswith(("- ", "* ", "+ ")):
            p.text = "• "
            add_markdown_text(p, line[2:], 19)
        elif re.match(r"^\d+[.)]\s", line):
            match = re.match(r"^(\d+[.)])\s+(.*)", line)
            p.text = match.group(1) + " "
            add_markdown_text(p, match.group(2), 19)
        elif line.startswith("> "):
            add_markdown_text(p, line[2:], 18)
            p.font.italic = True
        else:
            add_markdown_text(p, line, 19)

        p.space_after = Pt(10)


def markdown_to_pptx(input_file, output_file=None):
    input_path = Path(input_file)

    if not input_path.exists():
        raise FileNotFoundError(f"Файл не найден: {input_path}")

    # Читаем Markdown в UTF-8
    text = input_path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Разбираем документ на слайды по заголовкам # и ##
    sections = []
    title = input_path.stem
    current_title = title
    current_lines = []
    found_heading = False
    in_code = False

    for line in lines:
        if line.strip().startswith("```"):
            in_code = not in_code
            current_lines.append(line)
            continue

        match = None if in_code else re.match(r"^(#{1,2})\s+(.+)$", line)

        if match:
            found_heading = True

            if sections or current_lines or current_title != title:
                sections.append((current_title, current_lines))

            current_title = match.group(2).strip()
            current_lines = []
        else:
            current_lines.append(line)

    if current_title or current_lines:
        sections.append((current_title, current_lines))

    if not sections:
        sections = [(title, lines)]

    # Первый раздел оформляем как титульный слайд
    first_title, first_lines = sections[0]
    add_slide(prs, first_title, first_lines, title_slide=True)

    # Остальные разделы — обычные слайды
    for section_title, section_lines in sections[1:]:
        # Разбиваем слишком длинные разделы на части
        chunk_size = 8
        chunks = [
            section_lines[i:i + chunk_size]
            for i in range(0, len(section_lines), chunk_size)
        ] or [[]]

        for index, chunk in enumerate(chunks):
            slide_title = section_title
            if len(chunks) > 1:
                slide_title += f" ({index + 1})"

            add_slide(prs, slide_title, chunk)

    if output_file is None:
        output_file = input_path.with_suffix(".pptx")

    prs.save(output_file)
    print(f"Презентация сохранена: {output_file}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Использование: python maim.py файл.md [результат.pptx]")
        sys.exit(1)

    source = sys.argv[1]
    destination = sys.argv[2] if len(sys.argv) > 2 else None

    try:
        markdown_to_pptx(source, destination)
    except Exception as error:
        print(f"Ошибка: {error}")
        sys.exit(1)
